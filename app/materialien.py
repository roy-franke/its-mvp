"""L-01: Lernmaterialien global ablegen, mit Metadaten.

Ein Lernmaterial ist eine hochgeladene Datei, die jede Lektion verwenden kann.

Ablage (Entscheid E3): Die Originaldatei und der daraus gewonnene Text liegen
im Dateisystem, verwaltet mit filedepot (DEPOT), nie als BLOB in einer
Datenbank. Jede Datei bekommt einen Ordner `<ITS_MATERIAL_DIR>/depot/<uuid>/`
mit der Datei `file` und `metadata.json`. Der Katalog in SQLite (Tabelle
`materialien`, Entscheid E2) hält Metadaten, Datei-Metadaten, Speicherort und
den Stand der Indexierung (T-29). Die Material-ID verbindet Katalog, Dateien
und Embeddings in der Vektordatenbank.

Metadaten der Lehrperson: Name, mögliche Fächer, Niveau (Schätzung),
Beschreibung (höchstens 128 Zeichen), Tags. Automatisch: Dateiname, Grösse,
Format, Seitenzahl, letzte Änderung der Datei, Upload-Datum, Lehrperson.

Status: aktiv oder archiviert (L-03). Archivieren ist das Löschen; Datei,
Text und Embeddings bleiben erhalten, das Material ist nur nicht mehr wählbar.
"""

import hashlib
import io
import json
import logging
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
import time
import uuid
import zipfile
from pathlib import Path

from . import store

log = logging.getLogger("its.materialien")

BASE = Path(__file__).resolve().parent.parent
BESCHREIBUNG_MAX = 128
NAME_MAX = 120
WERT_MAX = 60
FORMATE = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "md": "text/markdown",
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
}
BILDFORMATE = ("png", "jpg", "jpeg")
LISTENFELDER = ("faecher", "tags")


class MaterialFehler(ValueError):
    """Fehler mit einer Meldung für die Lehrperson (HTTP 400)."""


class DublettenFehler(MaterialFehler):
    """Dieselbe Datei liegt schon in der Ablage (HTTP 409)."""

    def __init__(self, vorhanden: dict):
        super().__init__(f"Diese Datei liegt schon in der Ablage: «{vorhanden['name']}».")
        self.vorhanden = vorhanden


def material_dir() -> Path:
    return Path(os.getenv("ITS_MATERIAL_DIR") or BASE / "materialien")


def max_bytes() -> int:
    return int(float(os.getenv("ITS_MATERIAL_MAX_MB", "50")) * 1024 * 1024)


def _depot():
    from depot.io.local import LocalFileStorage
    pfad = material_dir() / "depot"
    pfad.mkdir(parents=True, exist_ok=True)
    return LocalFileStorage(str(pfad))


def speicherort(datei_id: str) -> str:
    """Pfad der Originaldatei, relativ zum Projektordner, falls möglich."""
    pfad = material_dir() / "depot" / datei_id / "file"
    try:
        return str(pfad.resolve().relative_to(BASE.resolve()))
    except ValueError:
        return str(pfad.resolve())


# ---------------------------------------------------------------- Datenbank

def _conn():
    c = sqlite3.connect(store.DB_PATH)
    c.row_factory = sqlite3.Row
    return c


def init_db():
    with _conn() as c:
        c.execute("""
            CREATE TABLE IF NOT EXISTS materialien (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                faecher TEXT NOT NULL DEFAULT '[]',
                niveau TEXT NOT NULL DEFAULT '',
                beschreibung TEXT NOT NULL DEFAULT '',
                tags TEXT NOT NULL DEFAULT '[]',
                dateiname TEXT NOT NULL,
                dateigroesse INTEGER NOT NULL,
                dateiformat TEXT NOT NULL,
                mime TEXT,
                seiten INTEGER,
                datei_geaendert_am REAL,
                hochgeladen_am REAL NOT NULL,
                hochgeladen_von TEXT NOT NULL DEFAULT '',
                datei_id TEXT NOT NULL,
                text_id TEXT,
                speicherort TEXT NOT NULL,
                sha256 TEXT NOT NULL,
                zeichen INTEGER NOT NULL DEFAULT 0,
                extraktion TEXT NOT NULL DEFAULT '{}',
                herkunft TEXT NOT NULL DEFAULT 'upload',
                status TEXT NOT NULL DEFAULT 'aktiv',
                archiviert_am REAL,
                geaendert_am REAL,
                index_status TEXT NOT NULL DEFAULT 'offen',
                index_info TEXT NOT NULL DEFAULT '{}'
            )
        """)
        c.execute("CREATE INDEX IF NOT EXISTS idx_materialien_sha ON materialien(sha256)")


def _row(r) -> dict:
    d = dict(r)
    for f in LISTENFELDER:
        d[f] = json.loads(d.get(f) or "[]")
    d["index_info"] = json.loads(d.get("index_info") or "{}")
    d["extraktion"] = json.loads(d.get("extraktion") or "{}")
    return d


def holen(mid: str) -> dict | None:
    with _conn() as c:
        r = c.execute("SELECT * FROM materialien WHERE id = ?", (mid,)).fetchone()
    return _row(r) if r else None


def alle(include_archived: bool = False) -> list[dict]:
    sql = "SELECT * FROM materialien"
    if not include_archived:
        sql += " WHERE status = 'aktiv'"
    with _conn() as c:
        rows = c.execute(sql + " ORDER BY hochgeladen_am DESC").fetchall()
    return [_row(r) for r in rows]


def setzen(mid: str, **felder):
    if not felder:
        return
    werte = []
    for k, v in felder.items():
        if k in LISTENFELDER or k in ("index_info", "extraktion"):
            v = json.dumps(v, ensure_ascii=False)
        werte.append(v)
    sets = ", ".join(f"{k} = ?" for k in felder)
    with _conn() as c:
        c.execute(f"UPDATE materialien SET {sets} WHERE id = ?", (*werte, mid))


# ---------------------------------------------------------------- Metadaten

def _bereinigen(text, maximum: int = WERT_MAX) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()[:maximum]


def _schluessel(text: str) -> str:
    return _bereinigen(text).casefold()


def werte() -> dict:
    """Bestehende Werte als Vorschläge (Autocompletion), aus allen Materialien."""
    felder = {"faecher": {}, "tags": {}, "niveaus": {}, "lehrpersonen": {}, "formate": {}}
    for m in alle(include_archived=True):
        for f in LISTENFELDER:
            for w in m[f]:
                felder[f].setdefault(_schluessel(w), w)
        if m["niveau"]:
            felder["niveaus"].setdefault(_schluessel(m["niveau"]), m["niveau"])
        if m["hochgeladen_von"]:
            felder["lehrpersonen"].setdefault(_schluessel(m["hochgeladen_von"]), m["hochgeladen_von"])
        felder["formate"].setdefault(m["dateiformat"], m["dateiformat"])
    return {k: sorted(v.values(), key=str.casefold) for k, v in felder.items()}


def kanonisch(wert: str, bestehende: list[str]) -> str:
    """Schreibweise eines bestehenden Werts übernehmen, wenn er sich nur in
    Gross-/Kleinschreibung oder Leerzeichen unterscheidet («tiere» -> «Tiere»)."""
    w = _bereinigen(wert)
    for b in bestehende:
        if _schluessel(b) == w.casefold():
            return b
    return w


def _liste(roh, bestehende: list[str]) -> list[str]:
    if isinstance(roh, str):
        try:
            geladen = json.loads(roh)
            roh = geladen if isinstance(geladen, list) else roh.split(",")
        except json.JSONDecodeError:
            roh = roh.split(",")
    out: list[str] = []
    for x in roh or []:
        w = kanonisch(x, bestehende + out)
        if w and _schluessel(w) not in {_schluessel(o) for o in out}:
            out.append(w)
    return out[:20]


def metadaten_pruefen(daten: dict, dateiname: str = "") -> dict:
    """Prüft und vereinheitlicht die Metadaten der Lehrperson."""
    vorhanden = werte()
    name = _bereinigen(daten.get("name"), NAME_MAX) or _bereinigen(Path(dateiname).stem, NAME_MAX)
    if not name:
        raise MaterialFehler("Name fehlt")
    beschreibung = re.sub(r"\s+", " ", str(daten.get("beschreibung") or "")).strip()
    if len(beschreibung) > BESCHREIBUNG_MAX:
        raise MaterialFehler(f"Die Beschreibung ist {len(beschreibung)} Zeichen lang, "
                             f"erlaubt sind höchstens {BESCHREIBUNG_MAX}.")
    return {
        "name": name,
        "faecher": _liste(daten.get("faecher"), vorhanden["faecher"]),
        "niveau": kanonisch(daten.get("niveau") or "", vorhanden["niveaus"]),
        "beschreibung": beschreibung,
        "tags": _liste(daten.get("tags"), vorhanden["tags"]),
    }


# ---------------------------------------------------------------- Extraktion

def dateiformat(dateiname: str) -> str:
    return Path(dateiname or "").suffix.lower().lstrip(".")


def _pdf_datum(text) -> float | None:
    """PDF-Datum «D:20260912143000+02'00'» in Sekunden."""
    m = re.match(r"D?:?(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?(\d{2})?", str(text or ""))
    if not m:
        return None
    import datetime
    teile = [int(x) if x else d for x, d in zip(m.groups(), (0, 1, 1, 0, 0, 0))]
    try:
        return datetime.datetime(*teile, tzinfo=datetime.timezone.utc).timestamp()
    except ValueError:
        return None


def _pdf(daten: bytes) -> dict:
    try:
        import fitz            # PyMuPDF, optional: bessere Extraktion (AGPL)
        doc = fitz.open(stream=daten, filetype="pdf")
        seiten = [{"seite": i + 1, "text": p.get_text("text") or ""} for i, p in enumerate(doc)]
        meta = doc.metadata or {}
        return {"seiten": seiten, "seitenzahl": len(seiten), "geaendert": _pdf_datum(meta.get("modDate")),
                "extraktor": "pymupdf"}
    except ImportError:
        pass
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(daten))
    seiten = [{"seite": i + 1, "text": p.extract_text() or ""} for i, p in enumerate(reader.pages)]
    geaendert = None
    try:
        geaendert = _pdf_datum((reader.metadata or {}).get("/ModDate"))
    except Exception:          # beschädigte Metadaten sind kein Grund zum Abbruch
        pass
    return {"seiten": seiten, "seitenzahl": len(seiten), "geaendert": geaendert, "extraktor": "pypdf"}


_UEBERSCHRIFT = re.compile(r"^(heading|überschrift|title|titel)\s*(\d)?", re.IGNORECASE)


def _docx(daten: bytes) -> dict:
    import docx
    d = docx.Document(io.BytesIO(daten))
    zeilen = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        m = _UEBERSCHRIFT.match(p.style.name if p.style is not None else "")
        if m:
            text = "#" * min(int(m.group(2) or 1), 4) + " " + text
        zeilen.append(text)
    for t in d.tables:
        for r in t.rows:
            zellen = [c.text.strip() for c in r.cells if c.text.strip()]
            if zellen:
                zeilen.append(" | ".join(dict.fromkeys(zellen)))
    seitenzahl = None
    try:
        with zipfile.ZipFile(io.BytesIO(daten)) as z:
            app = z.read("docProps/app.xml").decode("utf-8", "replace")
        m = re.search(r"<Pages>(\d+)</Pages>", app)
        seitenzahl = int(m.group(1)) if m else None
    except (KeyError, zipfile.BadZipFile):
        pass
    geaendert = d.core_properties.modified.timestamp() if d.core_properties.modified else None
    return {"seiten": [{"seite": None, "text": "\n\n".join(zeilen)}], "seitenzahl": seitenzahl,
            "geaendert": geaendert, "extraktor": "python-docx"}


def _text(daten: bytes) -> dict:
    for kodierung in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = daten.decode(kodierung)
            break
        except UnicodeDecodeError:
            continue
    return {"seiten": [{"seite": None, "text": text}], "seitenzahl": None, "geaendert": None,
            "extraktor": "text"}


def ocr_verfuegbar() -> bool:
    return shutil.which(os.getenv("ITS_TESSERACT", "tesseract")) is not None


def _bild(daten: bytes, fmt: str) -> dict:
    """Bilder: Text per OCR, falls Tesseract installiert ist, sonst kein Text."""
    out = {"seiten": [{"seite": 1, "text": ""}], "seitenzahl": 1, "geaendert": None, "extraktor": "keiner"}
    if not ocr_verfuegbar():
        out["hinweis"] = "Kein Text: Für Bilder braucht es Tesseract (OCR) auf dem Server."
        return out
    programm = os.getenv("ITS_TESSERACT", "tesseract")
    sprachen = os.getenv("ITS_TESSERACT_SPRACHEN", "deu+eng")
    with tempfile.TemporaryDirectory() as tmp:
        bild = Path(tmp) / f"bild.{fmt}"
        bild.write_bytes(daten)
        for lang in (sprachen, "eng"):
            r = subprocess.run([programm, str(bild), "stdout", "-l", lang], capture_output=True, timeout=120)
            if r.returncode == 0:
                out["seiten"][0]["text"] = r.stdout.decode("utf-8", "replace")
                out["extraktor"] = f"tesseract ({lang})"
                return out
    out["hinweis"] = "Die Texterkennung (OCR) ist gescheitert."
    return out


def extrahieren(dateiname: str, daten: bytes) -> dict:
    """Text seitenweise, Seitenzahl und Änderungsdatum aus einer Datei."""
    fmt = dateiformat(dateiname)
    if fmt not in FORMATE:
        raise MaterialFehler("Unterstützte Formate: " + ", ".join(f".{f}" for f in FORMATE))
    try:
        if fmt == "pdf":
            out = _pdf(daten)
        elif fmt == "docx":
            out = _docx(daten)
        elif fmt in BILDFORMATE:
            out = _bild(daten, fmt)
        else:
            out = _text(daten)
    except MaterialFehler:
        raise
    except Exception as e:
        raise MaterialFehler(f"Die Datei liess sich nicht lesen ({type(e).__name__}).") from e
    for s in out["seiten"]:
        s["text"] = re.sub(r"[ \t]+\n", "\n", re.sub(r"\n{3,}", "\n\n", s["text"] or "")).strip()
    out["text"] = "\n\n".join(s["text"] for s in out["seiten"] if s["text"])
    return out


# ---------------------------------------------------------------- Anlegen

def anlegen(dateiname: str, daten: bytes, metadaten: dict, hochgeladen_von: str = "",
            datei_geaendert_am: float | None = None, trotz_dublette: bool = False,
            herkunft: str = "upload") -> dict:
    """Legt ein Lernmaterial an: Datei und Text ins Depot, Metadaten in den Katalog."""
    dateiname = Path(dateiname or "").name
    if not dateiname:
        raise MaterialFehler("Dateiname fehlt")
    if len(daten) > max_bytes():
        raise MaterialFehler(f"Datei zu gross (höchstens {max_bytes() // (1024 * 1024)} MB)")
    if not daten:
        raise MaterialFehler("Die Datei ist leer")
    meta = metadaten_pruefen(metadaten, dateiname)
    sha = hashlib.sha256(daten).hexdigest()
    if not trotz_dublette:
        with _conn() as c:
            r = c.execute("SELECT * FROM materialien WHERE sha256 = ? AND status = 'aktiv' LIMIT 1", (sha,)).fetchone()
        if r:
            raise DublettenFehler(_row(r))
    fmt = dateiformat(dateiname)
    ext = extrahieren(dateiname, daten)
    if fmt not in BILDFORMATE and len(ext["text"]) < 50:
        raise MaterialFehler("Aus dieser Datei liess sich kaum Text extrahieren. "
                             "Ist es ein gescanntes PDF ohne Textebene?")
    depot = _depot()
    datei_id = depot.create(daten, filename=dateiname, content_type=FORMATE[fmt])
    text_json = json.dumps({"seiten": ext["seiten"], "extraktor": ext["extraktor"]}, ensure_ascii=False)
    text_id = depot.create(text_json.encode("utf-8"), filename="text.json", content_type="application/json")
    mid = "mat-" + uuid.uuid4().hex[:12]
    jetzt = time.time()
    geaendert = datei_geaendert_am if datei_geaendert_am else ext.get("geaendert")
    with _conn() as c:
        c.execute(
            "INSERT INTO materialien (id, name, faecher, niveau, beschreibung, tags, dateiname, dateigroesse, "
            "dateiformat, mime, seiten, datei_geaendert_am, hochgeladen_am, hochgeladen_von, datei_id, text_id, "
            "speicherort, sha256, zeichen, extraktion, herkunft, status, geaendert_am, index_status, index_info) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'aktiv', ?, ?, '{}')",
            (mid, meta["name"], json.dumps(meta["faecher"], ensure_ascii=False), meta["niveau"],
             meta["beschreibung"], json.dumps(meta["tags"], ensure_ascii=False), dateiname, len(daten), fmt,
             FORMATE[fmt], ext["seitenzahl"], geaendert, jetzt, _bereinigen(hochgeladen_von, 80) or "Lehrperson",
             datei_id, text_id, speicherort(datei_id), sha, len(ext["text"]),
             json.dumps({"extraktor": ext["extraktor"], "hinweis": ext.get("hinweis", "")}, ensure_ascii=False),
             herkunft, jetzt, "offen" if ext["text"] else "ohne_text"))
    log.info("Lernmaterial angelegt: %s (%s, %d Zeichen)", mid, dateiname, len(ext["text"]))
    return holen(mid)


def metadaten_aendern(mid: str, daten: dict) -> dict:
    m = holen(mid)
    if not m:
        raise KeyError(mid)
    meta = metadaten_pruefen({**{k: m[k] for k in ("name", "faecher", "niveau", "beschreibung", "tags")},
                              **{k: v for k, v in daten.items() if v is not None}}, m["dateiname"])
    setzen(mid, **meta, geaendert_am=time.time())
    return holen(mid)


# ---------------------------------------------------------------- Inhalte

def datei(mid: str) -> tuple[bytes, dict]:
    m = holen(mid)
    if not m:
        raise KeyError(mid)
    gespeichert = _depot().get(m["datei_id"])
    try:
        return gespeichert.read(), m
    finally:
        gespeichert.close()


def seiten(mid: str) -> list[dict]:
    """Der extrahierte Text, seitenweise: [{"seite": 1|None, "text": "..."}]."""
    m = holen(mid)
    if not m or not m.get("text_id"):
        return []
    gespeichert = _depot().get(m["text_id"])
    try:
        return json.loads(gespeichert.read().decode("utf-8")).get("seiten", [])
    finally:
        gespeichert.close()


def text(mid: str) -> str:
    return "\n\n".join(s["text"] for s in seiten(mid) if s.get("text"))


# ---------------------------------------------------------------- Status (L-03)

def archivieren(mid: str) -> dict:
    setzen(mid, status="archiviert", archiviert_am=time.time())
    log.info("Lernmaterial archiviert: %s", mid)
    return holen(mid)


def wiederherstellen(mid: str) -> dict:
    setzen(mid, status="aktiv", archiviert_am=None, geaendert_am=time.time())
    log.info("Lernmaterial wiederhergestellt: %s", mid)
    return holen(mid)


# ---------------------------------------------------------------- Filter (L-02)

def filtern(materialien: list[dict], q: str = "", fach: str = "", niveau: str = "", tag: str = "",
            fmt: str = "", lehrperson: str = "", von: float | None = None, bis: float | None = None) -> list[dict]:
    """Filter über Metadaten und Datei-Metadaten. Leere Filter wirken nicht."""
    def gleich(a, b):
        return _schluessel(a) == _schluessel(b)
    woerter = [w for w in _schluessel(q).split(" ") if w]
    out = []
    for m in materialien:
        if woerter:
            heu = " ".join([m["name"], m["beschreibung"], m["dateiname"], " ".join(m["tags"]),
                            " ".join(m["faecher"]), m["niveau"]]).casefold()
            if not all(w in heu for w in woerter):
                continue
        if fach and not any(gleich(f, fach) for f in m["faecher"]):
            continue
        if tag and not any(gleich(t, tag) for t in m["tags"]):
            continue
        if niveau and not gleich(m["niveau"], niveau):
            continue
        if fmt and m["dateiformat"] != fmt.lower().lstrip("."):
            continue
        if lehrperson and not gleich(m["hochgeladen_von"], lehrperson):
            continue
        if von and m["hochgeladen_am"] < von:
            continue
        if bis and m["hochgeladen_am"] > bis:
            continue
        out.append(m)
    return out


def oeffentlich(m: dict) -> dict:
    """Katalogeintrag für die Oberfläche (ohne interne Depot-IDs)."""
    return {k: v for k, v in m.items() if k not in ("datei_id", "text_id", "sha256")}


def quellen_block(m: dict, text_inhalt: str | None = None) -> str:
    """Material als Block im Lektionsmaterial (gleiches Format wie bisher, N-01)."""
    return f"### Quelle: {m['name']}\n\n{text_inhalt if text_inhalt is not None else text(m['id'])}"
